#!/usr/bin/env python3
"""
build_deck.py -- builds presentation/project-overview.pptx (16:9, 6 slides, native editable shapes + chart).

    python presentation/prepare_assets.py      # once: extracts the two real photos from data/
    python presentation/build_deck.py          # builds the deck
    powershell presentation/render_deck.ps1    # optional: PDF + PNG per slide via PowerPoint (Windows)

Requirements: python-pptx, pillow (see presentation/requirements.txt).

Where the numbers come from (nothing is invented):
  * Problem facts (10 Hz, 4 m between frames, < 1 m lateral target, hidden test may differ in light/speed):
        docs/*.pdf (organisers' problem statement and intro)
  * Night-session facts (1194 s, 11,943 samples, top speed): measured from data/training_cache.h5
        (presentation/assets/facts.json, written by prepare_assets.py)
  * Before/after milestone (24.508 -> 13.688 m, 98 % along-track, ...): README.md and
        brainstorming_and_diagnostics.md (git tags rmse-24.508, rmse-13.688). These were measured on the EARLIER
        pipeline version, and the deck says so.
  * "88 % of the squared error in the last 9 s": chats/astra.md (debugging log).
  * How the current pipeline works: Scripts/*.py as they are in the working tree.
  * Results of the CURRENT pipeline, runtime and hardware are NOT recorded anywhere in the repository. They are read
        from presentation/results_to_fill.json. While a value is null, the slide shows a clearly marked
        amber "fill in" placeholder instead of a number.
"""
import json
import os

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, 'assets')
OUT_PPTX = os.path.join(HERE, 'project-overview.pptx')
RESULTS_JSON = os.path.join(HERE, 'results_to_fill.json')

# ------------------------------------------------------------------------------------------------------------------
# Design tokens (reference deck: white page, Arial, left-aligned titles, grey secondary text, pastel diagram blocks)
# ------------------------------------------------------------------------------------------------------------------
FONT = 'Arial'
INK, GREY, GREY_L, LINE_C, PANEL, WHITE = '1B1F24', '5F6B7A', '9AA5B1', 'D5DBE3', 'F2F4F7', 'FFFFFF'
GREY_M = '6B7787'   # smallest text (footers, sources): darker than GREY_L for legibility
BLUE, BLUE_P = '2F6FDE', 'E3EDFC'
ORANGE, ORANGE_P = 'C4560F', 'FCEBDD'
GREEN, GREEN_P = '1E8A5A', 'DDF2E7'
RED, RED_P = 'B83227', 'FBE4E1'
AMBER, AMBER_P, AMBER_L = '8A5A00', 'FFF4D6', 'E0A800'

SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.6
CONTENT_W = SLIDE_W - 2 * MARGIN          # 12.133 in
N_SLIDES = 6

# ------------------------------------------------------------------------------------------------------------------
# Numbers that ARE recorded in the repository (README.md, brainstorming_and_diagnostics.md; tags rmse-24.508/13.688)
# ------------------------------------------------------------------------------------------------------------------
BASE = dict(rmse=24.508, median=9.003, p95=56.861, long=24.302, lat=3.172, yaw=2.159, within1=12.8)
TUNED = dict(rmse=13.688, median=1.547, p95=20.004, long=13.308, lat=3.201, yaw=2.368, within1=32.8)
LONG_SHARE = BASE['long'] ** 2 / BASE['rmse'] ** 2           # 0.983
RMSE_DROP = 1.0 - TUNED['rmse'] / BASE['rmse']               # 0.44


def load_facts():
    with open(os.path.join(ASSETS, 'facts.json')) as fh:
        return json.load(fh)


def load_results():
    """Values the repository does not contain yet. Missing file -> template with nulls is written."""
    template = {
        '_help': 'Fill these in after you have run the final pipeline; null = placeholder is shown on the slides. '
                 'Sources: tools/visualize.py metrics.txt (validation), the [TIMING] line printed by '
                 'Scripts/4_infer.py (ms per frame).',
        'team_name': None,
        'repo_url': None,
        'val_position_rmse_m': None,
        'val_lateral_rmse_m': None,
        'ms_per_frame_mean': None,
        'hardware': None,
    }
    if not os.path.exists(RESULTS_JSON):
        with open(RESULTS_JSON, 'w') as fh:
            json.dump(template, fh, indent=2)
    with open(RESULTS_JSON) as fh:
        data = json.load(fh)
    return {k: data.get(k) for k in template if not k.startswith('_')}


# ------------------------------------------------------------------------------------------------------------------
# Low-level helpers
# ------------------------------------------------------------------------------------------------------------------
def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def run_fmt(run, size, bold=False, color=INK, italic=False):
    f = run.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = rgb(color)


def P(content, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, after=0, italic=False):
    """A paragraph spec. `content` is a string or a list of (text, {overrides}) runs."""
    runs = [(content, {})] if isinstance(content, str) else content
    return dict(runs=runs, size=size, bold=bold, color=color, align=align, after=after, italic=italic)


def fill_text(tf, paras, anchor=MSO_ANCHOR.TOP, margins=(0.1, 0.06, 0.1, 0.06)):
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE                 # never shrink text to fit
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(m) for m in margins]
    tf.vertical_anchor = anchor
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p['align']
        para.space_after = Pt(p['after'])
        for text, over in p['runs']:
            r = para.add_run()
            r.text = text
            run_fmt(r, over.get('size', p['size']), over.get('bold', p['bold']), over.get('color', p['color']),
                    over.get('italic', p['italic']))


def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, name=None, margins=(0.0, 0.0, 0.0, 0.0)):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    fill_text(tb.text_frame, paras, anchor, margins)
    if name:
        tb.name = name
    return tb


def box(slide, x, y, w, h, fill=None, line=None, line_w=1.25, dash=False, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        radius=0.08, paras=None, anchor=MSO_ANCHOR.MIDDLE, name=None, margins=(0.12, 0.06, 0.12, 0.06)):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(line_w)
        if dash:
            s.line.dash_style = MSO_LINE.DASH
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    if paras:
        fill_text(s.text_frame, paras, anchor, margins)
    if name:
        s.name = name
    return s


def arrow(slide, x1, y1, x2, y2, color=GREY, width=2.25, dash=False, head=True, both=False, name=None):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = MSO_LINE.DASH
    ln = c.line._get_or_add_ln()
    for tag, on in (('a:headEnd', both), ('a:tailEnd', head)):
        if on:
            el = etree.SubElement(ln, qn(tag))
            el.set('type', 'triangle')
            el.set('w', 'med')
            el.set('len', 'med')
    if name:
        c.name = name
    return c


def picture(slide, path, x, y, w, alt, name=None, border=LINE_C):
    pic = slide.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w))
    pic._element.nvPicPr.cNvPr.set('descr', alt)
    pic.line.color.rgb = rgb(border)
    pic.line.width = Pt(0.75)
    if name:
        pic.name = name
    return pic


def pill(slide, x, y, w, h, label, fill=INK, color=WHITE, size=14, name=None):
    return box(slide, x, y, w, h, fill=fill, radius=0.5, name=name, margins=(0.08, 0.02, 0.08, 0.02),
               paras=[P(label, size=size, bold=True, color=color, align=PP_ALIGN.CENTER)])


def placeholder_chip(slide, x, y, w, h, label, size=14, name=None):
    """Amber dashed chip = information that the repository does not contain yet."""
    return box(slide, x, y, w, h, fill=AMBER_P, line=AMBER_L, line_w=1.5, dash=True, radius=0.2, name=name,
               margins=(0.08, 0.02, 0.08, 0.02),
               paras=[P(label, size=size, bold=True, color=AMBER, align=PP_ALIGN.CENTER)])


# ------------------------------------------------------------------------------------------------------------------
# Slide scaffolding
# ------------------------------------------------------------------------------------------------------------------
def new_slide(prs, number, kicker, title, title_size=30):
    slide = prs.slides.add_slide(prs.slide_layouts[5])                 # "Title Only": keeps real titles for outline view
    t = slide.shapes.title
    t.left, t.top, t.width, t.height = Inches(MARGIN), Inches(0.62), Inches(CONTENT_W), Inches(0.75)
    fill_text(t.text_frame, [P(title, size=title_size, bold=True, color=INK)], MSO_ANCHOR.TOP, (0, 0, 0, 0))
    t.name = 'Title'
    text(slide, MARGIN, 0.3, 9.5, 0.3, [P(kicker, size=14, bold=True, color=BLUE)], name='Kicker')
    footer(slide, number)
    return slide


def footer(slide, number):
    text(slide, MARGIN, 7.08, 9, 0.28,
         [P('Camera-only localization of a race car  ·  Constructor Labs challenge', size=12, color=GREY_M)],
         name='Footer')
    text(slide, SLIDE_W - MARGIN - 1.5, 7.08, 1.5, 0.28,
         [P(f'{number} / {N_SLIDES}', size=12, color=GREY_M, align=PP_ALIGN.RIGHT)], name='Slide number')


def set_notes(slide, note):
    slide.notes_slide.notes_text_frame.text = note.strip()


def val_or_fill(value, fmt, label):
    """Formatted value if the results file has it, else a visible [FILL IN: ...] marker (used in speaker notes)."""
    return fmt.format(value) if value is not None else f'[FILL IN: {label}]'


# ------------------------------------------------------------------------------------------------------------------
# Slides
# ------------------------------------------------------------------------------------------------------------------
def slide_problem(prs, facts, res):
    s = new_slide(prs, 1, 'NATIONAL AI HACKATHON 2026  ·  CONSTRUCTOR LABS CHALLENGE', 'Where is the car?', 44)
    s.shapes.title.top = Inches(0.62)
    s.shapes.title.height = Inches(0.85)
    text(s, MARGIN, 1.5, CONTENT_W, 0.4,
         [P('Cameras only. No GNSS, no LiDAR, no peeking at future frames.', size=22, color=GREY)], name='Subtitle')

    if res['team_name']:
        text(s, 9.4, 0.3, SLIDE_W - MARGIN - 9.4, 0.3,
             [P(res['team_name'], size=14, bold=True, color=GREY, align=PP_ALIGN.RIGHT)], name='Team name')
    else:
        placeholder_chip(s, SLIDE_W - MARGIN - 2.4, 0.27, 2.4, 0.36, 'Team name: fill in', name='Placeholder: team name')

    pw = 5.95
    ph = pw * facts['night_size'][1] / facts['night_size'][0]
    py = 2.15
    px2 = SLIDE_W - MARGIN - pw
    picture(s, os.path.join(ASSETS, 'night_fast.jpg'), MARGIN, py, pw,
            'Front-left camera frame from the night training session: dark track, floodlights and a gantry',
            name='Photo: night training frame')
    picture(s, os.path.join(ASSETS, 'day_sample.jpg'), px2, py, pw,
            'Front-left camera frame from the daylight sample recording: grandstand and control tower in daylight',
            name='Photo: daylight sample frame')
    pill(s, MARGIN + 0.15, py + 0.15, 2.35, 0.36, f'Night  ·  {facts["night_pick_speed"]:.0f} m/s', name='Badge: night')
    pill(s, px2 + 0.15, py + 0.15, 2.35, 0.36, 'Daylight sample', fill=BLUE, name='Badge: daylight')
    cap_y = py + ph + 0.08
    text(s, MARGIN, cap_y, pw, 0.35,
         [P(f'Training: {facts["duration_s"]:.0f} s of night driving', size=16, color=GREY)], name='Caption: night')
    text(s, px2, cap_y, pw, 0.35,
         [P('Hidden test: may be faster, in different light', size=16, color=GREY)], name='Caption: test')

    cards = [('10 Hz', 'one pose for every camera frame'),
             ('4 m', 'driven between two frames at 40 m/s'),
             ('< 1 m', 'sideways accuracy needed to hold a racing line')]
    cw, gap, cy, ch = 3.85, (CONTENT_W - 3 * 3.85) / 2, 5.78, 1.1
    for i, (big, small) in enumerate(cards):
        cx = MARGIN + i * (cw + gap)
        box(s, cx, cy, cw, ch, fill=PANEL, name=f'Card {i + 1}')
        text(s, cx + 0.2, cy, 1.45, ch, [P(big, size=34, bold=True, color=BLUE)], MSO_ANCHOR.MIDDLE, name=f'Card {i + 1} number')
        text(s, cx + 1.7, cy, cw - 1.85, ch, [P(small, size=14, color=INK)], MSO_ANCHOR.MIDDLE, name=f'Card {i + 1} label')

    set_notes(s, """
Picture a race car at forty metres per second, at night. Between two camera frames it has travelled four metres. We have to say where it is, to under a metre sideways, using cameras only: no GPS, no lidar. And the estimate may only use images up to now, which is what causal means, exactly as on the real car.
Our training recording is a night session, on the left. But the hidden test may be faster and in different light, like the daylight sample on the right. So the problem is really two problems: locate precisely, and do it in light we have never labelled.
Next: our idea in one picture.
""")


def slide_idea(prs):
    s = new_slide(prs, 2, 'OUR IDEA', 'Recognise the place. Track the motion. Fuse both.')

    xs = dict(c1=(0.6, 1.9), c2=(2.95, 2.2), c3=(5.6, 2.2), f=(8.25, 2.3), o=(11.0, 1.733))
    ya, yb, bh = 1.85, 3.85, 1.45

    text(s, MARGIN, 1.5, 7.0, 0.3, [P('WHAT THE CAMERAS SEE', size=14, bold=True, color=BLUE)], name='Lane label: see')
    text(s, MARGIN, 3.45, 7.0, 0.3, [P('HOW THE CAR MOVED', size=14, bold=True, color=ORANGE)], name='Lane label: move')

    def block(key, y, title, sub, fill, line, dash=False, name=None):
        x, w = xs[key]
        return box(s, x, y, w, bh, fill=fill, line=line, dash=dash, name=name,
                   paras=[P(title, size=16, bold=True, color=INK, after=3), P(sub, size=14, color=INK)])

    block('c1', ya, '4 cameras', 'Front and side views, cropped to the track', BLUE_P, BLUE, name='Block: cameras')
    block('c2', ya, 'Place network', 'ResNet-18 turns each view into a 256-number fingerprint', BLUE_P, BLUE, name='Block: place network')
    block('c3', ya, 'Map lookup', 'Finds the most similar training views near our estimate', BLUE_P, BLUE, name='Block: map lookup')
    block('c1', yb, 'Car telemetry', 'Speed and turn rate, as the rules allow', ORANGE_P, ORANGE, name='Block: telemetry')
    block('c2', yb, 'Dead reckoning', 'Adds up distance and turning since the last frame', ORANGE_P, ORANGE, name='Block: dead reckoning')
    block('c3', yb, 'Cross-check', 'Motion learned from camera pairs is the backup', ORANGE_P, ORANGE,
          dash=True, name='Block: cross-check')

    fx, fw = xs['f']
    box(s, fx, ya, fw, yb + bh - ya, fill=GREEN_P, line=GREEN, name='Block: Kalman filter',
        paras=[P('Kalman filter', size=20, bold=True, after=8),
               P('Predicts with the motion, then corrects with the map match.', size=16, after=8),
               P('Each is weighted by how much it is trusted.', size=16)])
    ox, ow = xs['o']
    cy = (ya + yb + bh) / 2
    box(s, ox, cy - 0.65, ow, 1.3, fill=INK, name='Block: output pose',
        paras=[P('Pose', size=20, bold=True, color=WHITE, align=PP_ALIGN.CENTER, after=3),
               P('x, y, heading, 10 times a second', size=14, color=WHITE, align=PP_ALIGN.CENTER)])

    def row_arrows(y_mid, color):
        for a, b in (('c1', 'c2'), ('c2', 'c3')):
            arrow(s, xs[a][0] + xs[a][1] + 0.04, y_mid, xs[b][0] - 0.04, y_mid, color=color)
        arrow(s, xs['c3'][0] + xs['c3'][1] + 0.04, y_mid, fx - 0.04, y_mid, color=color)

    row_arrows(ya + bh / 2, BLUE)
    row_arrows(yb + bh / 2, ORANGE)
    arrow(s, fx + fw + 0.04, cy, ox - 0.04, cy, color=GREEN, width=3)

    defs = [('Fingerprint', 'A list of 256 numbers describing a view. Similar views give similar numbers.'),
            ('Map', 'Every training frame stored with its fingerprint and its known position.'),
            ('Kalman filter', 'A standard way to blend a prediction with a measurement, trusting each by its confidence.')]
    dw = (CONTENT_W - 2 * 0.25) / 3
    for i, (term, desc) in enumerate(defs):
        box(s, MARGIN + i * (dw + 0.25), 5.55, dw, 1.3, fill=PANEL, name=f'Definition: {term}', anchor=MSO_ANCHOR.TOP,
            margins=(0.15, 0.1, 0.15, 0.06), paras=[P('In plain words: ' + term, size=14, bold=True, color=GREY, after=3),
                                                   P(desc, size=14, color=INK)])

    set_notes(s, """
Here is the whole system. Top lane, the cameras. A neural network turns each view into a fingerprint, a list of 256 numbers; similar views give similar fingerprints. We built a map from the training laps: each frame's fingerprint with its known position. Looking one up tells us roughly where we are.
Bottom lane, the car's own speed and turn rate, which the rules allow. They say how far we moved, whatever the light. Motion learned from camera pairs is the backup.
A Kalman filter blends both, trusting each by its confidence, and outputs x, y and heading ten times a second.
Next: why each part is there.
""")


def slide_answers(prs):
    s = new_slide(prs, 3, 'HOW IT HANDLES THE HARD PARTS', 'Each known difficulty gets a specific answer')

    cols = dict(d=(MARGIN, 3.8), a=(5.05, 5.4), t=(10.75, 1.983))
    for key, label in (('d', 'THE DIFFICULTY'), ('a', 'OUR ANSWER'), ('t', 'WHEN IT RUNS')):
        x, w = cols[key]
        text(s, x, 1.48, w, 0.3, [P(label, size=14, bold=True, color=GREY)], name=f'Column header: {label.lower()}')

    rows = [
        ('Trained at night, tested in unknown light',
         'Lighting augmentation, robust inputs',
         'Training images get random lighting, shadows, glare and haze; every camera image is contrast-normalised.',
         'train + run time'),
        ('Look-alike straights',
         'Learn places, then trust by similarity',
         'Training pairs show one place on different laps. Weak matches count less; only nearby matches are averaged.',
         'training + filter'),
        ('4 m between frames, blur, cameras out of sync',
         'Strictly causal camera sync',
         'Each camera supplies its newest image stamped no later than now. A learned gate weights the four views.',
         'every frame'),
        ('Drift, and getting lost',
         'Motion from the car, global re-lock',
         'Calibrated speed and turn rate drive the prediction. A filter over the whole lap finds the start and re-locks.',
         'every frame'),
    ]
    y0, rh, step = 1.85, 1.05, 1.15
    for i, (diff, head, detail, when) in enumerate(rows):
        y = y0 + i * step
        dx, dw = cols['d']
        ax, aw = cols['a']
        tx, tw = cols['t']
        box(s, dx, y, dw, rh, fill=RED_P, name=f'Difficulty {i + 1}',
            paras=[P(diff, size=16, bold=True, color=INK)])
        arrow(s, dx + dw + 0.08, y + rh / 2, ax - 0.08, y + rh / 2, color=GREY)
        box(s, ax, y, aw, rh, fill=BLUE_P, name=f'Answer {i + 1}', anchor=MSO_ANCHOR.MIDDLE,
            paras=[P(head, size=16, bold=True, color=INK, after=2), P(detail, size=14, color=INK)])
        pill(s, tx, y + rh / 2 - 0.2, tw, 0.4, when, fill=WHITE, color=GREY, size=14, name=f'Tag {i + 1}')
        s.shapes[-1].line.color.rgb = rgb(LINE_C)
        s.shapes[-1].line.width = Pt(1)

    text(s, MARGIN, 6.6, CONTENT_W, 0.3,
         [P('Training-only parts add no cost at run time. These are design choices; the next slide shows the one measured result we have.',
            size=14, color=GREY)], name='Note')

    set_notes(s, """
Each part answers a difficulty the organisers named.
Night versus unknown light: training images get random lighting, shadows, glare and haze, and every camera image is contrast-normalised, in training and at run time.
Look-alike straights: we train on same-place, different-lap pairs and trust a weak match less.
Timing: a tenth of a second is four metres, so each camera contributes only its newest image no later than now.
Drift and getting lost: motion comes from the car's calibrated speed and turn rate. We do not assume a start pose: a filter over the whole lap compares each view with the entire map and commits once one place clearly wins; the same filter re-locks the car when it is lost.
Now the evidence.
""")


def slide_results(prs, res):
    s = new_slide(prs, 4, 'EVIDENCE', 'Measure first: a filter fix cut the error by 44%')

    text(s, MARGIN, 1.5, 6.6, 0.3, [P('Position error in metres (lower is better)', size=16, bold=True, color=GREY)],
         name='Chart heading')
    cd = CategoryChartData()
    cd.categories = ['Total (RMSE)', 'Along the track', 'Across the track']
    cd.add_series('Before filter tuning', (BASE['rmse'], BASE['long'], BASE['lat']))
    cd.add_series('After filter tuning', (TUNED['rmse'], TUNED['long'], TUNED['lat']))
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(MARGIN), Inches(1.85), Inches(6.7), Inches(4.45), cd)
    gf.name = 'Chart: error before and after tuning'
    ch = gf.chart
    ch.font.size = Pt(14)
    ch.font.name = FONT
    ch.font.color.rgb = rgb(INK)
    ch.has_title = False
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(14)
    plot = ch.plots[0]
    plot.gap_width = 55
    plot.overlap = -8
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = '0.0'
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(16)
    dl.font.bold = True
    for series, color in zip(plot.series, (GREY_L, GREEN)):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = rgb(color)
    va = ch.value_axis
    va.visible = False
    va.has_major_gridlines = False
    va.minimum_scale = 0
    va.maximum_scale = 29
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(14)
    ca.format.line.color.rgb = rgb(LINE_C)
    ca.has_major_gridlines = False

    text(s, MARGIN, 6.38, 6.7, 0.6,
         [P('Held-out validation, earlier pipeline version (git tags rmse-24.508 and rmse-13.688). Sources: README.md, brainstorming_and_diagnostics.md.',
            size=12, color=GREY_M)], name='Chart source')

    rx, rw = 7.75, SLIDE_W - MARGIN - 7.75
    text(s, rx, 1.45, 2.3, 1.1, [P(f'\u2212{RMSE_DROP * 100:.0f}%', size=60, bold=True, color=GREEN)], MSO_ANCHOR.MIDDLE,
         name='Big number')
    text(s, rx + 2.4, 1.5, rw - 2.4, 1.0,
         [P('position error, with no retraining: two filter trust settings changed', size=16, color=INK)],
         MSO_ANCHOR.MIDDLE, name='Big number label')
    stats = [
        (f'{LONG_SHARE * 100:.0f}%', 'of the squared error was along the track'),
        (f'{BASE["median"]:.1f} to {TUNED["median"]:.1f} m', 'median error'),
        (f'{BASE["within1"]:.1f}% to {TUNED["within1"]:.1f}%', 'of frames within 1 m of the truth'),
    ]
    for i, (big, small) in enumerate(stats):
        y = 2.7 + i * 0.62
        text(s, rx, y, 2.3, 0.55, [P(big, size=18, bold=True, color=INK)], MSO_ANCHOR.MIDDLE, name=f'Stat {i + 1} number')
        text(s, rx + 2.4, y, rw - 2.4, 0.55, [P(small, size=14, color=GREY)], MSO_ANCHOR.MIDDLE, name=f'Stat {i + 1} label')

    # Final pipeline: not recorded in the repository -> placeholders until results_to_fill.json is filled in
    items = [('Validation position RMSE', res['val_position_rmse_m'], '{:.2f} m'),
             ('Across-track (lateral) RMSE', res['val_lateral_rmse_m'], '{:.2f} m'),
             ('Time per frame', res['ms_per_frame_mean'], '{:.1f} ms'),
             ('Hardware', res['hardware'], '{}')]
    complete = all(v is not None for _, v, _ in items)
    cy, chh = 4.7, 2.2
    box(s, rx, cy, rw, chh, fill=GREEN_P if complete else AMBER_P, line=GREEN if complete else AMBER_L, line_w=1.5,
        dash=not complete, name='Card: final pipeline results')
    text(s, rx + 0.2, cy + 0.1, rw - 0.4, 0.35,
         [P('Final pipeline' if complete else 'Final pipeline: fill in', size=16, bold=True,
            color=GREEN if complete else AMBER)], name='Card heading')
    for i, (label, value, fmt) in enumerate(items):
        y = cy + 0.55 + i * 0.4
        text(s, rx + 0.2, y, 2.9, 0.36, [P(label, size=14, color=INK)], MSO_ANCHOR.MIDDLE, name=f'Final {i + 1} label')
        shown = fmt.format(value) if value is not None else 'fill in'
        text(s, rx + 3.1, y, rw - 3.3, 0.36,
             [P(shown, size=14, bold=True, color=INK if value is not None else AMBER, align=PP_ALIGN.RIGHT)],
             MSO_ANCHOR.MIDDLE, name=f'Final {i + 1} value')

    rmse = val_or_fill(res['val_position_rmse_m'], '{:.1f} metres', 'final validation position RMSE')
    ms = val_or_fill(res['ms_per_frame_mean'], '{:.0f} milliseconds', 'ms per frame')
    hw = val_or_fill(res['hardware'], '{}', 'hardware')
    set_notes(s, f"""
Before adding anything, we measured. An offline diagnostic tool, never part of the submission, showed 98 percent of the squared error was along the track. The filter trusted its motion too much and the map match too little.
Re-weighting two settings, with no retraining, cut position error from 24.5 to 13.7 metres, minus 44 percent. The median fell from 9.0 to 1.5 metres. Error across the track stayed near 3 metres, as predicted.
Those figures are from the earlier pipeline. For the final one: validation error {rmse}, and {ms} per frame on {hw}.
Next: where it still fails.
""")


def slide_failures(prs):
    s = new_slide(prs, 5, 'FAILURE CASES', 'Where it still fails, and why')

    cw, gap, y0, chh = 3.85, (CONTENT_W - 3 * 3.85) / 2, 1.6, 5.1
    cards = [
        ('Look-alike straights',
         'Along-track error dominated: 98% of the early squared error. After tuning, 5% of frames are still more than 20 m off.',
         'Response: trust the map match more. Next: match short sequences of views.'),
        ('Places not on the map',
         'An earlier validation run ended parked on an unmapped spot: 88% of its squared error came from the last 9 seconds.',
         'Response: track map with the pit corridor, standstill handling, whole-lap re-lock. Not yet re-measured.'),
        ('Unseen daylight',
         'Lighting augmentation is a proxy, not a measurement. The daylight sample has no ground-truth poses to score against.',
         'Response: lighting augmentation, robust normalisation, a match test relative to the lap, speed from the car.'),
    ]
    for i, (title, body, response) in enumerate(cards):
        x = MARGIN + i * (cw + gap)
        box(s, x, y0, cw, chh, fill=PANEL, name=f'Failure card {i + 1}')
        text(s, x + 0.2, y0 + 0.12, cw - 0.4, 0.4, [P(title, size=20, bold=True, color=RED)], name=f'Failure {i + 1} title')
        dx, dy, dw, dh = x + 0.2, y0 + 0.65, cw - 0.4, 1.6
        box(s, dx, dy, dw, dh, fill=WHITE, name=f'Failure {i + 1} diagram frame')
        if i == 0:
            diagram_straight(s, dx, dy, dw, dh)
        elif i == 1:
            diagram_unmapped(s, dx, dy, dw, dh)
        else:
            diagram_light(s, dx, dy, dw, dh)
        text(s, x + 0.2, y0 + 2.45, cw - 0.4, 1.5, [P(body, size=16, color=INK)], name=f'Failure {i + 1} text')
        text(s, x + 0.2, y0 + 3.95, cw - 0.4, 1.05, [P(response, size=14, color=GREY)], name=f'Failure {i + 1} response')

    text(s, MARGIN, 6.78, CONTENT_W, 0.26,
         [P('Sources: README.md and brainstorming log (98%, 20 m); chats/astra.md debugging log (88%, 9 s); problem statement (daylight test).',
            size=12, color=GREY_M)], name='Sources')

    set_notes(s, """
Three failure cases, honestly.
One, look-alike straights. Along-track error is still our weak spot: even after tuning, five percent of frames are still more than twenty metres off.
Two, places the map has never seen. In an earlier validation run the car ended parked on an unmapped spot, and nearly nine-tenths of the squared error came from the last nine seconds. We have since added a track map with the pit corridor, standstill handling and a whole-lap re-lock, but we have not re-measured that yet.
Three, daylight. Our lighting augmentation is a proxy, not a measurement, and the daylight sample has no ground-truth poses to score against.
Lessons next.
""")


def diagram_straight(s, x, y, w, h):
    top, bot = y + 0.2, y + h - 0.2
    arrow(s, x + 0.15, top, x + w - 0.15, top, color=GREY_L, width=2, head=False, name='Road edge top')
    arrow(s, x + 0.15, bot, x + w - 0.15, bot, color=GREY_L, width=2, head=False, name='Road edge bottom')
    mid = y + h / 2
    d = 0.26
    t = box(s, x + 0.45, mid - d / 2, d, d, fill=INK, shape=MSO_SHAPE.OVAL, name='True position')
    e = box(s, x + w - 0.95, mid - d / 2, d, d, fill=ORANGE, shape=MSO_SHAPE.OVAL, name='Estimated position')
    arrow(s, x + 0.45 + d + 0.04, mid, x + w - 0.95 - 0.04, mid, color=RED, width=2, both=True, name='Along-track error')
    text(s, x + 0.2, y + 0.26, 1.2, 0.3, [P('true', size=14, bold=True)], name='Label: true')
    text(s, x + w - 1.3, y + 0.26, 1.2, 0.3, [P('estimate', size=14, bold=True, color=ORANGE, align=PP_ALIGN.RIGHT)],
         name='Label: estimate')
    text(s, x + 0.2, mid + 0.2, w - 0.4, 0.3, [P('along-track error', size=14, color=RED, align=PP_ALIGN.CENTER)],
         name='Label: error')


def diagram_unmapped(s, x, y, w, h):
    mid = y + h / 2 - 0.1
    split = x + w * 0.45
    arrow(s, x + 0.15, mid, split, mid, color=GREEN, width=5, head=False, name='Mapped stretch')
    arrow(s, split, mid, x + w - 0.15, mid, color=RED, width=5, head=False, name='Unmapped stretch')
    box(s, x + w - 0.55, mid - 0.13, 0.26, 0.26, fill=INK, shape=MSO_SHAPE.OVAL, name='Parked car')
    text(s, x + 0.15, mid + 0.22, split - x - 0.1, 0.3, [P('in the map', size=14, bold=True, color=GREEN)], name='Label: mapped')
    text(s, split + 0.05, mid + 0.22, x + w - split - 0.2, 0.3,
         [P('not in the map', size=14, bold=True, color=RED, align=PP_ALIGN.RIGHT)], name='Label: unmapped')
    text(s, x + 0.15, y + 0.12, w - 0.3, 0.3, [P('parked here at the end', size=14, color=GREY, align=PP_ALIGN.RIGHT)],
         name='Label: parked')


def diagram_light(s, x, y, w, h):
    sq, sy = 0.62, y + 0.2
    gap = (w - 0.4 - 3 * sq) / 2
    xs = [x + 0.2 + i * (sq + gap) for i in range(3)]
    box(s, xs[0], sy, sq, sq, fill='1E2630', name='Night sample', radius=0.12)
    box(s, xs[1], sy, sq, sq, fill='8FA3B8', name='Augmented night sample', radius=0.12)
    box(s, xs[2], sy, sq, sq, fill='DCE4EC', line=AMBER_L, line_w=1.5, dash=True, name='Real daylight sample', radius=0.12,
        paras=[P('?', size=20, bold=True, color=AMBER, align=PP_ALIGN.CENTER)], margins=(0, 0, 0, 0))
    arrow(s, xs[0] + sq + 0.05, sy + sq / 2, xs[1] - 0.05, sy + sq / 2, color=GREY, width=2, name='Augmented')
    arrow(s, xs[1] + sq + 0.05, sy + sq / 2, xs[2] - 0.05, sy + sq / 2, color=AMBER_L, width=2, dash=True, name='Unverified')
    ly = sy + sq + 0.1
    for i, (label, col) in enumerate((('night', INK), ('augmented', GREY), ('real day', AMBER))):
        text(s, xs[i] - 0.35, ly, sq + 0.7, 0.3, [P(label, size=14, bold=True, color=col, align=PP_ALIGN.CENTER)],
             name=f'Label: {label}')


def slide_lessons(prs, res):
    s = new_slide(prs, 6, 'LESSONS', 'Lessons learned and what comes next')

    lessons = [
        ('Diagnose before you train', 'An offline tool that may use ground truth showed the fault was the filter, not the network.'),
        ('Calibrated trust beats a bigger model', 'Two re-weighted settings gave minus 44%, with no retraining.'),
        ('Use every signal the rules allow', "The car's speed and turn rate do not depend on light."),
    ]
    for i, (head, sub) in enumerate(lessons):
        y = 1.65 + i * 1.4
        box(s, MARGIN, y + 0.12, 0.6, 0.6, fill=INK, shape=MSO_SHAPE.OVAL, name=f'Lesson {i + 1} badge', margins=(0, 0, 0, 0),
            paras=[P(str(i + 1), size=20, bold=True, color=WHITE, align=PP_ALIGN.CENTER)])
        text(s, MARGIN + 0.85, y, 6.1, 1.1,
             [P(head, size=20, bold=True, after=3), P(sub, size=16, color=GREY)], name=f'Lesson {i + 1}')

    nx, nw = 8.1, SLIDE_W - MARGIN - 8.1
    text(s, nx, 1.5, nw, 0.3, [P('NEXT', size=14, bold=True, color=BLUE)], name='Next heading')
    steps = [('Short sequences of views', 'instead of one frame at a time'),
             ('Geometric refinement', 'once the place is known: the path to sub-metre'),
             ('Pick the best checkpoint', 'by trajectory error, not the last epoch')]
    for i, (head, sub) in enumerate(steps):
        y = 1.9 + i * 1.4
        box(s, nx, y, nw, 1.0, fill=BLUE_P, line=BLUE, dash=True, name=f'Next step {i + 1}',
            paras=[P(head, size=16, bold=True, after=2), P(sub, size=14)])
        if i < 2:
            arrow(s, nx + nw / 2, y + 1.0 + 0.03, nx + nw / 2, y + 1.4 - 0.03, color=BLUE, width=2.25)

    sy = 5.85
    box(s, MARGIN, sy, CONTENT_W, 1.05, fill=PANEL, name='Declared and reproducible', anchor=MSO_ANCHOR.MIDDLE,
        margins=(0.2, 0.06, 2.8, 0.06),
        paras=[P([('Declared pre-trained model: ', {'bold': True}), ('ImageNet ResNet-18 (torchvision). ', {}),
                  ('No external datasets; the daylight sample appears only as a slide photo.', {})], size=14, after=2),
               P('Code, README and requirements.txt are in the repository.', size=14, color=GREY)])
    if res['repo_url']:
        text(s, SLIDE_W - MARGIN - 2.6, sy, 2.45, 1.05, [P(res['repo_url'], size=14, bold=True, color=BLUE, align=PP_ALIGN.RIGHT)],
             MSO_ANCHOR.MIDDLE, name='Repository link')
    else:
        placeholder_chip(s, SLIDE_W - MARGIN - 2.5, sy + 0.33, 2.35, 0.4, 'Repo link: fill in', name='Placeholder: repository link')

    set_notes(s, """
Three lessons. Diagnose before you train: one offline tool told us the fault was the filter, not the network. Calibrated trust beats a bigger model: two settings gave minus forty-four percent. And use every signal the rules allow: the car's speed and turn rate do not depend on light.
Next, we would match short sequences of views instead of single frames, and refine the position geometrically once the place is known, which is the path to sub-metre accuracy.
Everything is in the repository with a README and requirements. The one pre-trained model is ImageNet ResNet-18, declared here; there is no external data, and the daylight sample is only the photo on the first slide.
Thank you.
""")


# ------------------------------------------------------------------------------------------------------------------
def main():
    facts = load_facts()
    res = load_results()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
    prs.core_properties.title = 'Camera-only race car localization'
    prs.core_properties.subject = 'National AI Hackathon 2026, Constructor Labs challenge'

    slide_problem(prs, facts, res)
    slide_idea(prs)
    slide_answers(prs)
    slide_results(prs, res)
    slide_failures(prs)
    slide_lessons(prs, res)
    prs.save(OUT_PPTX)

    words = 0
    for i, sl in enumerate(prs.slides, 1):
        n = len([w for w in sl.notes_slide.notes_text_frame.text.split() if not w.startswith('[FILL')])
        words += n
        print(f'slide {i}: {n:3d} words of notes (~{n / 145 * 60:.0f} s at 145 wpm)')
    print(f'total {words} words, about {words / 145:.1f} min at 145 wpm (limit 5 min)')
    print(f'wrote {OUT_PPTX}')


if __name__ == '__main__':
    main()
